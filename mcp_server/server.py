"""Local MCP server for the HPA gene-specificity screening pipeline.

Run directly for stdio, or with ``mcp dev mcp_server/server.py`` for Inspector.
The original command-line scripts remain the source of all scientific computations.
"""
from __future__ import annotations

import csv
import json
import logging
import os
import re
import subprocess
import sys
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Annotated, Any, Literal

from mcp.server import MCPServer
from mcp.server.mcpserver.exceptions import ToolError
from mcp.types import ToolAnnotations
from openpyxl import load_workbook
from pydantic import BaseModel, Field

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CONFIG = ROOT / "assets" / "config.example.json"
RUNS_DIR = Path(
    os.environ.get("HPA_MCP_RUNS_DIR", ROOT / "mcp_runs")
).expanduser().resolve()
LOGGER = logging.getLogger("hpa-gene-specificity-mcp")
RUN_ID_PATTERN = re.compile(r"[0-9a-f]{12}")


class ArtifactPaths(BaseModel):
    """Absolute paths to the durable files produced by a screening run."""

    expression_workbook: str
    selected_genes_tsv: str
    screening_summary_json: str
    merged_workbook: str | None = None
    plots_directory: str | None = None
    validation_step1_json: str | None = None
    validation_step2_json: str | None = None


class ScreeningResult(BaseModel):
    """A compact result page plus links to the complete local artifacts."""

    run_id: str
    status: Literal["completed"]
    summary: dict[str, Any]
    selected_genes: list[dict[str, str]]
    selected_gene_count: int
    returned_gene_count: int
    offset: int
    limit: int
    truncated: bool
    artifacts: ArtifactPaths


class CellTypeCatalog(BaseModel):
    """Cell-type names accepted by a particular HPA expression workbook."""

    expression_workbook: str
    cell_type_count: int
    cell_types: list[str]


mcp = MCPServer(
    "HPA Gene Specificity Screen",
    instructions=(
        "Screen gene lists for cell-type-specific expression with Human Protein "
        "Atlas single-cell RNA-seq data. Use run_specificity_screen for a new "
        "gene list, screen_existing_workbook to reuse an HPA workbook, and "
        "get_screening_results to paginate through selected genes. The screening "
        "criterion is the peak cell type; tissue expression is annotation only."
    ),
)


def _input_path(value: str, label: str) -> Path:
    path = Path(value).expanduser()
    if not path.is_absolute():
        path = ROOT / path
    path = path.resolve()
    if not path.is_file():
        raise ToolError(f"{label} does not exist or is not a file: {path}")
    return path


def _config_path(value: str | None) -> Path:
    return DEFAULT_CONFIG if value is None else _input_path(value, "Config file")


def _load_config(path: Path) -> dict[str, Any]:
    try:
        config = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ToolError(f"Cannot read valid JSON config {path}: {exc}") from exc
    groups = config.get("specificity_groups")
    if not isinstance(groups, dict) or not groups:
        raise ToolError("Config must define a non-empty 'specificity_groups' object.")
    for name, members in groups.items():
        if (
            not isinstance(name, str)
            or not name.strip()
            or not isinstance(members, list)
            or not members
            or not all(isinstance(member, str) and member.strip() for member in members)
        ):
            raise ToolError(
                "Each specificity group must have a non-empty name and a "
                "non-empty list of HPA cell-type names."
            )
    return config


def _new_run_dir() -> tuple[str, Path]:
    RUNS_DIR.mkdir(parents=True, exist_ok=True)
    run_id = uuid.uuid4().hex[:12]
    run_dir = RUNS_DIR / run_id
    run_dir.mkdir(exist_ok=False)
    return run_id, run_dir


def _run_dir(run_id: str) -> Path:
    if not RUN_ID_PATTERN.fullmatch(run_id):
        raise ToolError("Invalid run_id; use the 12-character ID returned by a screening run.")
    path = (RUNS_DIR / run_id).resolve()
    if path.parent != RUNS_DIR or not path.is_dir():
        raise ToolError(f"Screening run not found: {run_id}")
    return path


def _output_tail(value: str, limit: int = 6000) -> str:
    value = value.strip()
    return value if len(value) <= limit else "...\n" + value[-limit:]


def _run_command(command: list[str], label: str) -> None:
    """Run an existing CLI step without invoking a shell."""
    LOGGER.info("Starting %s", label)
    result = subprocess.run(
        command,
        cwd=ROOT,
        text=True,
        capture_output=True,
        check=False,
    )
    if result.stdout.strip():
        LOGGER.info("%s output:\n%s", label, result.stdout.strip())
    if result.stderr.strip():
        LOGGER.info("%s diagnostics:\n%s", label, result.stderr.strip())
    if result.returncode != 0:
        raise ToolError(
            f"{label} failed with exit code {result.returncode}.\n"
            f"stdout:\n{_output_tail(result.stdout) or '(empty)'}\n"
            f"stderr:\n{_output_tail(result.stderr) or '(empty)'}"
        )


def _workbook_cell_types(workbook_path: Path) -> list[str]:
    try:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        if "RNA_Cell_Types" not in workbook.sheetnames:
            workbook.close()
            raise ToolError(
                f"Workbook has no 'RNA_Cell_Types' sheet: {workbook_path}"
            )
        worksheet = workbook["RNA_Cell_Types"]
        header = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
        workbook.close()
    except ToolError:
        raise
    except Exception as exc:
        raise ToolError(f"Cannot read HPA workbook {workbook_path}: {exc}") from exc
    return [str(value) for value in header[2:] if value not in (None, "")]


def _validate_cell_types(workbook_path: Path, config: dict[str, Any]) -> None:
    available = set(_workbook_cell_types(workbook_path))
    requested = {
        member
        for members in config["specificity_groups"].values()
        for member in members
    }
    unknown = sorted(requested - available, key=str.casefold)
    if unknown:
        raise ToolError(
            "These configured cell types are absent from the HPA workbook: "
            + ", ".join(unknown)
            + ". Call list_available_cell_types with this workbook to see valid names."
        )


def _write_manifest(
    run_dir: Path,
    run_id: str,
    expression_workbook: Path,
    make_plots: bool,
) -> None:
    manifest = {
        "run_id": run_id,
        "status": "completed",
        "completed_at": datetime.now(timezone.utc).isoformat(),
        "expression_workbook": str(expression_workbook),
        "make_plots": make_plots,
    }
    (run_dir / "run_manifest.json").write_text(
        json.dumps(manifest, indent=2), encoding="utf-8"
    )


def _optional_file(path: Path) -> str | None:
    return str(path) if path.is_file() else None


def _read_results(run_id: str, offset: int, limit: int) -> ScreeningResult:
    run_dir = _run_dir(run_id)
    summary_path = run_dir / "screening_summary.json"
    selected_path = run_dir / "selected_genes.tsv"
    manifest_path = run_dir / "run_manifest.json"
    try:
        summary = json.loads(summary_path.read_text(encoding="utf-8"))
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        with selected_path.open(encoding="utf-8", newline="") as handle:
            rows = [dict(row) for row in csv.DictReader(handle, delimiter="\t")]
    except (OSError, json.JSONDecodeError, csv.Error) as exc:
        raise ToolError(f"Results for run {run_id} are incomplete or unreadable: {exc}") from exc

    page = rows[offset:offset + limit]
    plots = run_dir / "gene_plots"
    artifacts = ArtifactPaths(
        expression_workbook=manifest["expression_workbook"],
        selected_genes_tsv=str(selected_path),
        screening_summary_json=str(summary_path),
        merged_workbook=_optional_file(run_dir / "HPA_RNA_with_GEP.xlsx"),
        plots_directory=str(plots) if plots.is_dir() else None,
        validation_step1_json=_optional_file(run_dir / "validation_step1.json"),
        validation_step2_json=_optional_file(run_dir / "validation_step2.json"),
    )
    return ScreeningResult(
        run_id=run_id,
        status="completed",
        summary=summary,
        selected_genes=page,
        selected_gene_count=len(rows),
        returned_gene_count=len(page),
        offset=offset,
        limit=limit,
        truncated=offset + len(page) < len(rows),
        artifacts=artifacts,
    )


@mcp.tool(
    title="Run HPA gene-specificity screen",
    annotations=ToolAnnotations(
        read_only_hint=False,
        destructive_hint=False,
        idempotent_hint=False,
        open_world_hint=True,
    ),
)
def run_specificity_screen(
    gene_list_path: Annotated[
        str,
        Field(description="CSV, TSV, or XLSX gene list with a gene, symbol, or query_gene column."),
    ],
    config_path: Annotated[
        str | None,
        Field(description="Screening JSON config; omit to use assets/config.example.json."),
    ] = None,
    make_plots: Annotated[
        bool,
        Field(description="Generate per-gene cell-type and tissue plots and the merged workbook."),
    ] = True,
    workers: Annotated[
        int,
        Field(ge=1, le=16, description="Parallel plot workers; ignored when make_plots is false."),
    ] = 4,
) -> ScreeningResult:
    """Download HPA expression data and screen a gene list for cell-type specificity.

    This accesses the public HPA API and may take several minutes. Every invocation
    creates an isolated run directory. The returned run_id retrieves later result pages.
    """
    gene_list = _input_path(gene_list_path, "Gene-list file")
    config_file = _config_path(config_path)
    config = _load_config(config_file)
    run_id, run_dir = _new_run_dir()
    workbook = run_dir / "HPA_RNA_expression.xlsx"

    _run_command([
        sys.executable,
        str(ROOT / "scripts" / "step1_fetch_hpa.py"),
        "--gene-list", str(gene_list),
        "--output-dir", str(run_dir),
        "--config", str(config_file),
    ], "HPA download")
    _validate_cell_types(workbook, config)

    if make_plots:
        _run_command([
            sys.executable,
            str(ROOT / "scripts" / "step2_plot_and_merge.py"),
            "--gene-list", str(gene_list),
            "--hpa-excel", str(workbook),
            "--output-dir", str(run_dir),
            "--config", str(config_file),
            "--workers", str(workers),
        ], "Plot and workbook merge")

    _run_command([
        sys.executable,
        str(ROOT / "scripts" / "screen_specificity.py"),
        "--hpa-excel", str(workbook),
        "--config", str(config_file),
        "--out-dir", str(run_dir),
    ], "Specificity screening")
    _write_manifest(run_dir, run_id, workbook, make_plots)
    return _read_results(run_id, offset=0, limit=50)


@mcp.tool(
    title="Screen an existing HPA workbook",
    annotations=ToolAnnotations(
        read_only_hint=False,
        destructive_hint=False,
        idempotent_hint=False,
        open_world_hint=False,
    ),
)
def screen_existing_workbook(
    hpa_excel_path: Annotated[
        str,
        Field(description="Existing HPA_RNA_expression.xlsx produced by step 1."),
    ],
    config_path: Annotated[
        str | None,
        Field(description="Screening JSON config; omit to use assets/config.example.json."),
    ] = None,
) -> ScreeningResult:
    """Apply a new specificity configuration without downloading HPA data again."""
    workbook = _input_path(hpa_excel_path, "HPA workbook")
    config = _load_config(_config_path(config_path))
    _validate_cell_types(workbook, config)
    run_id, run_dir = _new_run_dir()
    _run_command([
        sys.executable,
        str(ROOT / "scripts" / "screen_specificity.py"),
        "--hpa-excel", str(workbook),
        "--config", str(_config_path(config_path)),
        "--out-dir", str(run_dir),
    ], "Specificity screening")
    _write_manifest(run_dir, run_id, workbook, make_plots=False)
    return _read_results(run_id, offset=0, limit=50)


@mcp.tool(
    title="Get screening results",
    annotations=ToolAnnotations(read_only_hint=True, open_world_hint=False),
)
def get_screening_results(
    run_id: Annotated[str, Field(description="The 12-character ID returned by a screening run.")],
    offset: Annotated[int, Field(ge=0, description="Zero-based selected-gene offset.")] = 0,
    limit: Annotated[int, Field(ge=1, le=500, description="Maximum selected genes to return.")] = 50,
) -> ScreeningResult:
    """Read a screening summary and a bounded page of selected genes."""
    return _read_results(run_id, offset=offset, limit=limit)


@mcp.tool(
    title="List HPA cell types",
    annotations=ToolAnnotations(read_only_hint=True, open_world_hint=False),
)
def list_available_cell_types(
    hpa_excel_path: Annotated[
        str,
        Field(description="HPA workbook whose RNA_Cell_Types headers should be listed."),
    ],
) -> CellTypeCatalog:
    """List exact cell-type names accepted in specificity_groups for this workbook."""
    workbook = _input_path(hpa_excel_path, "HPA workbook")
    cell_types = _workbook_cell_types(workbook)
    if not cell_types:
        raise ToolError("The RNA_Cell_Types sheet contains no cell-type columns.")
    return CellTypeCatalog(
        expression_workbook=str(workbook),
        cell_type_count=len(cell_types),
        cell_types=cell_types,
    )


if __name__ == "__main__":
    mcp.run(transport="stdio")
